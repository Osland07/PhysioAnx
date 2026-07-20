from PySide6.QtWidgets import QTableWidgetItem, QWidget, QHBoxLayout, QPushButton, QLabel
from PySide6.QtCore import Qt
import qtawesome as qta

class ReportController:
    def __init__(self, view, main_window):
        self.view = view
        self.main_window = main_window
        
        # Simulasi load data
        self.load_dummy_data()

    def load_dummy_data(self):
        dummy_sesi = [
            ("19 Juli 2026, 09:15", "2406-001", "Bpk. Budi", "Severe", "45 Tahun", "Laki-laki", "170 cm", "75 kg", "Jl. Merdeka No. 1, Jakarta"),
            ("18 Juli 2026, 14:30", "2406-002", "Ibu Siti", "Moderate", "38 Tahun", "Perempuan", "160 cm", "60 kg", "Jl. Sudirman No. 12, Bandung"),
            ("17 Juli 2026, 10:00", "2406-003", "Sdr. Andi", "Mild", "25 Tahun", "Laki-laki", "175 cm", "68 kg", "Jl. Diponegoro No. 8, Surabaya")
        ]
        
        self.view.table_sesi.setRowCount(len(dummy_sesi))
        for row, data in enumerate(dummy_sesi):
            item_tgl = QTableWidgetItem(data[0])
            item_tgl.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            self.view.table_sesi.setItem(row, 0, item_tgl)
            
            item_id = QTableWidgetItem(data[1])
            item_id.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            self.view.table_sesi.setItem(row, 1, item_id)
            
            item_nama = QTableWidgetItem(data[2])
            item_nama.setTextAlignment(Qt.AlignVCenter | Qt.AlignLeft)
            self.view.table_sesi.setItem(row, 2, item_nama)
            
            # Badge Indikasi Kecemasan
            lbl_indikasi = QLabel(data[3])
            lbl_indikasi.setAlignment(Qt.AlignCenter)
            if data[3] == "Severe":
                lbl_indikasi.setStyleSheet("background-color: #E53E3E; color: #FFF; padding: 4px; border-radius: 4px; font-weight: bold;")
            elif data[3] == "Moderate":
                lbl_indikasi.setStyleSheet("background-color: #DD6B20; color: #FFF; padding: 4px; border-radius: 4px; font-weight: bold;")
            else:
                lbl_indikasi.setStyleSheet("background-color: #38A169; color: #FFF; padding: 4px; border-radius: 4px; font-weight: bold;")
                
            ind_widget = QWidget()
            ind_layout = QHBoxLayout(ind_widget)
            ind_layout.setContentsMargins(5, 5, 5, 5)
            ind_layout.addWidget(lbl_indikasi)
            self.view.table_sesi.setCellWidget(row, 3, ind_widget)
            
            # Kolom Aksi
            action_widget = QWidget()
            action_layout = QHBoxLayout(action_widget)
            action_layout.setContentsMargins(5, 5, 5, 5)
            action_layout.setSpacing(8)
            
            btn_detail = QPushButton(" Detail")
            btn_detail.setIcon(qta.icon('fa5s.eye', color='white'))
            btn_detail.setStyleSheet("color: white; font-weight: bold; background: #3182CE; border-radius: 4px; padding: 4px 8px; font-size: 11px;")
            btn_detail.setCursor(Qt.PointingHandCursor)
            
            btn_pdf = QPushButton(" PDF")
            btn_pdf.setIcon(qta.icon('fa5s.file-pdf', color='white'))
            btn_pdf.setStyleSheet("color: white; font-weight: bold; background: #ED2224; border-radius: 4px; padding: 4px 8px; font-size: 11px;")
            btn_pdf.setCursor(Qt.PointingHandCursor)
            
            btn_excel = QPushButton(" Excel")
            btn_excel.setIcon(qta.icon('fa5s.file-excel', color='white'))
            btn_excel.setStyleSheet("color: white; font-weight: bold; background: #107C41; border-radius: 4px; padding: 4px 8px; font-size: 11px;")
            btn_excel.setCursor(Qt.PointingHandCursor)
            
            btn_detail.clicked.connect(lambda checked=False, d=data: self.open_detail_replay(d))
            btn_pdf.clicked.connect(lambda checked=False, d=data: self.export_pdf(d))
            btn_excel.clicked.connect(lambda checked=False, d=data: self.export_excel(d))
            
            action_layout.addWidget(btn_detail)
            action_layout.addWidget(btn_pdf)
            action_layout.addWidget(btn_excel)
            
            self.view.table_sesi.setCellWidget(row, 4, action_widget)
            
        self.view.lbl_info.setText(f"Menampilkan 1 hingga {len(dummy_sesi)} dari {len(dummy_sesi)} riwayat sesi")

    def export_excel(self, data):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        import openpyxl
        
        # Bersihkan nama pasien untuk nama file
        safe_name = data[2].replace('.', '').replace(' ', '_')
        filepath, _ = QFileDialog.getSaveFileName(self.view, "Simpan Excel", f"Laporan_{safe_name}.xlsx", "Excel Files (*.xlsx)")
        if filepath:
            try:
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Laporan Sesi"
                
                # Header
                ws.append(["Tanggal Sesi", "ID Pasien", "Nama Pasien", "Indikasi Kecemasan"])
                # Data
                ws.append([data[0], data[1], data[2], data[3]])
                
                # Styling sederhana untuk Header
                from openpyxl.styles import Font
                for cell in ws[1]:
                    cell.font = Font(bold=True)
                    
                wb.save(filepath)
                QMessageBox.information(self.view, "Sukses", f"Data Excel berhasil diekspor ke:\n{filepath}")
            except Exception as e:
                QMessageBox.critical(self.view, "Error", f"Gagal menyimpan file:\n{str(e)}")

    def export_pdf(self, data):
        from PySide6.QtWidgets import QFileDialog, QMessageBox
        from PySide6.QtGui import QTextDocument, QPageSize
        from PySide6.QtPrintSupport import QPrinter
        from datetime import datetime

        safe_name = data[2].replace('.', '').replace(' ', '_')
        filepath, _ = QFileDialog.getSaveFileName(
            self.view, "Save PDF Report",
            f"PhysioReport_{safe_name}.pdf", "PDF Files (*.pdf)"
        )
        if not filepath:
            return

        try:
            printer = QPrinter(QPrinter.HighResolution)
            printer.setOutputFormat(QPrinter.PdfFormat)
            printer.setOutputFileName(filepath)
            printer.setPageSize(QPageSize(QPageSize.A4))

            sev = data[3]
            sev_colors = {
                "Severe":   ("#C53030", "#FFF5F5"),
                "Moderate": ("#C05621", "#FFFAF0"),
                "Mild":     ("#975A16", "#FFFFF0"),
                "Minimal":  ("#276749", "#F0FFF4"),
            }
            sev_fg, sev_bg = sev_colors.get(sev, ("#4A5568", "#EDF2F7"))
            
            now = datetime.now()
            months_id = ["", "Januari", "Februari", "Maret", "April", "Mei", "Juni", "Juli", "Agustus", "September", "Oktober", "November", "Desember"]
            tanggal_ttd = f"{now.day} {months_id[now.month]} {now.year}"
            printed_at = datetime.now().strftime("%d %B %Y, %H:%M")

            html = f"""
<!DOCTYPE html>
<html>
<head>
<meta charset="utf-8">
<style>
  * {{ box-sizing: border-box; }}
  body {{
    font-family: "Times New Roman", Times, serif;
    font-size: 11pt;
    color: #000;
    line-height: 1.3;
    background: #fff;
    padding: 10px 20px;
  }}

  /* ── KOP SURAT ── */
  .kop-surat {{
    text-align: center;
    border-bottom: 3px solid #000;
    padding-bottom: 5px;
    margin-bottom: 2px;
    line-height: 1.1;
  }}
  .kop-surat h1 {{
    font-size: 14pt;
    margin: 0;
    text-transform: uppercase;
  }}
  .kop-surat h2 {{
    font-size: 12pt;
    margin: 0;
  }}
  .kop-surat p {{
    font-size: 10pt;
    margin: 0;
  }}
  .kop-garis-bawah {{
    border-bottom: 1px solid #000;
    margin-bottom: 15px;
  }}

  /* ── JUDUL SURAT ── */
  .judul-surat {{
    text-align: center;
    font-weight: bold;
    font-size: 12pt;
    text-decoration: underline;
    margin-bottom: 15px;
  }}

  /* ── ISI SURAT ── */
  .content {{
    text-align: justify;
  }}
  .table-identitas {{
    width: 100%;
    margin-bottom: 10px;
    margin-left: 20px;
    border-collapse: collapse;
  }}
  .table-identitas td {{
    vertical-align: top;
    padding: 1px 3px;
  }}

  /* ── TABEL HASIL ── */
  .table-hasil {{
    width: 100%;
    border-collapse: collapse;
    margin: 10px 0;
  }}
  .table-hasil th, .table-hasil td {{
    border: 1px solid #000;
    padding: 5px;
    text-align: left;
  }}
  .table-hasil th {{
    text-align: center;
  }}

  /* ── TANDA TANGAN ── */
  .tanda-tangan {{
    width: 250px;
    float: right;
    text-align: center;
    margin-top: 20px;
  }}
  .tanda-tangan p {{
    margin: 0;
  }}
  .nama-terang {{
    font-weight: bold;
    text-decoration: underline;
    margin-top: 50px !important;
  }}
</style>
</head>
<body>

<!-- KOP SURAT -->
<div class="kop-surat">
  <h1>KLINIK FISIOTERAPI PHYSIOANX</h1>
  <h2>Pusat Pemantauan Fisiologis & Kecemasan</h2>
  <p>Jl. Kesehatan No. 123, Jakarta, Indonesia | Telp: (021) 555-1234</p>
</div>
<div class="kop-garis-bawah"></div>

<!-- JUDUL SURAT -->
<div class="judul-surat">SURAT HASIL PEMERIKSAAN FISIOLOGIS</div>

<!-- ISI SURAT -->
<div class="content">
  <p>Yang bertanda tangan di bawah ini, menerangkan bahwa pasien berikut:</p>
  
  <table class="table-identitas">
    <tr>
      <td width="20%">Nama</td>
      <td width="2%">:</td>
      <td width="78%"><b>{data[2]}</b></td>
    </tr>
    <tr>
      <td>No. Rekam Medis</td>
      <td>:</td>
      <td>{data[1]}</td>
    </tr>
    <tr>
      <td>Usia</td>
      <td>:</td>
      <td>{data[4]}</td>
    </tr>
    <tr>
      <td>Jenis Kelamin</td>
      <td>:</td>
      <td>{data[5]}</td>
    </tr>
    <tr>
      <td>Tinggi / Berat</td>
      <td>:</td>
      <td>{data[6]} / {data[7]}</td>
    </tr>
  </table>

  <p>Telah menjalani pemeriksaan pemantauan respons fisiologis pada tanggal <b>{data[0]}</b>. Berdasarkan hasil perekaman sensor biometrik selama sesi, didapatkan ringkasan parameter sebagai berikut:</p>

  <table class="table-hasil">
    <tr>
      <th width="30%">Parameter</th>
      <th width="20%">Rata-rata</th>
      <th width="20%">Puncak</th>
      <th width="30%">Keterangan</th>
    </tr>
    <tr>
      <td>Heart Rate (BPM)</td>
      <td>84.2 bpm</td>
      <td>112.5 bpm</td>
      <td>Reaktivitas kardiovaskular tinggi</td>
    </tr>
    <tr>
      <td>Skin Conductance (&mu;S)</td>
      <td>5.8 &mu;S</td>
      <td>9.1 &mu;S</td>
      <td>Aktivitas kelenjar keringat aktif</td>
    </tr>
    <tr>
      <td>Skin Temperature (&deg;C)</td>
      <td>31.5 &deg;C</td>
      <td>33.2 &deg;C</td>
      <td>Vasokonstriksi (menurun)</td>
    </tr>
  </table>

  <p><b>Kesimpulan Pemeriksaan:</b></p>
  <p>Berdasarkan analisis data fluktuasi <i>Heart Rate</i> dan lonjakan <i>Galvanic Skin Response</i> (GSR) di atas, pasien terindikasi mengalami tingkat kecemasan <b>{sev.upper()}</b>. Hasil ini menunjukkan adanya hiperaktivitas pada saraf simpatis selama sesi perekaman.</p>

  <p>Demikian surat hasil pemeriksaan ini dibuat untuk dapat dipergunakan sebagaimana mestinya.</p>
</div>

<!-- TANDA TANGAN -->
<div class="tanda-tangan">
  <p>Jakarta, {tanggal_ttd}</p>
  <p>Dokter / Fisioterapis Pemeriksa,</p>
  <p class="nama-terang">dr. Pemeriksa PhysioAnx, Sp.Kj</p>
  <p>NIP. 19801231 200501 1 001</p>
</div>

</body>
</html>
"""
            doc = QTextDocument()
            doc.setHtml(html)
            doc.print_(printer)

            QMessageBox.information(
                self.view, "Success",
                f"PDF report successfully exported to:\n{filepath}"
            )
        except Exception as e:
            QMessageBox.critical(self.view, "Error", f"Failed to save PDF:\n{str(e)}")

    def open_detail_replay(self, data):
        from components.replay_dialog import ReplayDialog
        dialog = ReplayDialog(data, self.main_window)
        dialog.exec()
